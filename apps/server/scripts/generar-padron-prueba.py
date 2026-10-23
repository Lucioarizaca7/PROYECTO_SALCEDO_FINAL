"""Genera libros Excel de prueba para "Importar padrón" (POST /api/v1/comerciantes/importar-padron).

Salida (datos-prueba/):
  - padron-prueba-50.xlsx: 50 comerciantes válidos en dos hojas (BLOQUE A y BLOQUE B).
  - padron-prueba-errores.xlsx: filas con errores intencionales para probar la validación.

Los datos son ficticios y reproducibles (semilla fija). Uso:
  python3 apps/server/scripts/generar-padron-prueba.py
"""
import random
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT_DIR = Path(__file__).resolve().parents[3] / 'datos-prueba'

# Encabezados reconocidos por los alias de apps/server/src/padronImport.js.
HEADER = ['N°', 'DNI/RUC', 'Titular', 'Código Puesto', 'Sector', 'Zona',
          'Tipo Comerciante', 'Rubro', 'Metraje m2', 'Estado', 'Fecha Feria']

NOMBRES = ['Yolanda', 'Rosa', 'Juana', 'Nidia', 'Maribel', 'Elena', 'Lourdes', 'Gladys',
           'Marina', 'Sonia', 'Felicitas', 'Teodora', 'Ruth', 'Hilda', 'Delia', 'Vilma',
           'Juan', 'Pedro', 'Mario', 'Rubén', 'Edwin', 'Wilber', 'Percy', 'Fredy',
           'Alberto', 'Raúl', 'Hugo', 'Néstor', 'Efraín', 'Roger', 'Germán', 'Santos']
APELLIDOS = ['Mamani', 'Quispe', 'Condori', 'Choque', 'Ticona', 'Apaza', 'Huanca', 'Cutipa',
             'Layme', 'Pari', 'Coaquira', 'Flores', 'Ramos', 'Chambi', 'Calsin', 'Ccama',
             'Turpo', 'Aguilar', 'Vilca', 'Machaca', 'Pacompia', 'Larico', 'Yupanqui', 'Atencio']

RUBROS = ['comida', 'ropa', 'zapatillas', 'otros']
TIPOS = ['Permanente', 'Ferial', 'Eventual']

# Cada bloque corresponde a una feria: martes 20 y miércoles 21 de octubre de 2026.
BLOQUES = [
    ('BLOQUE A', 'A', ['Mercado Central'], date(2026, 10, 20)),
    ('BLOQUE B', 'B', ['Parque Principal', 'Zona Norte'], date(2026, 10, 21)),
]


def ruc_persona_natural(dni):
    """RUC 10 + DNI + dígito verificador (módulo 11 de SUNAT)."""
    base = '10' + dni
    pesos = [5, 4, 3, 2, 7, 6, 5, 4, 3, 2]
    resto = sum(int(d) * p for d, p in zip(base, pesos)) % 11
    return base + str((11 - resto) % 10)


def generar_filas(rng, total_por_bloque=25):
    usados = set()
    filas = {}
    for hoja, letra, zonas, fecha in BLOQUES:
        filas[hoja] = []
        for i in range(1, total_por_bloque + 1):
            while True:
                dni = f'{rng.randint(1000000, 79999999):08d}'
                if dni not in usados:
                    usados.add(dni)
                    break
            documento = ruc_persona_natural(dni) if rng.random() < 0.2 else dni
            titular = (f'{rng.choice(NOMBRES)} {rng.choice(APELLIDOS)} {rng.choice(APELLIDOS)}').upper()
            fila = {
                'n': i,
                'documento': documento,
                'titular': titular,
                'codigo': f'PRB-{letra}-{i:03d}',
                'sector': f'Bloque {letra} - Fila {(i - 1) // 5 + 1}',
                'zona': zonas[(i - 1) % len(zonas)],
                'tipo': rng.choices(TIPOS, weights=[5, 4, 1])[0],
                'rubro': rng.choices(RUBROS, weights=[3, 3, 3, 1])[0],
                'metraje': rng.choice([2, 2.5, 3, 3.5, 4, 4.5, 6, 8, 9, 12]),
                'estado': 'Presente',
                'fecha': fecha,
            }
            filas[hoja].append(fila)

    # 6 ausentes (3 por bloque): se registra la falta de esa feria, que no se cobra; el puesto sigue asignado.
    for hoja in filas:
        for fila in rng.sample(filas[hoja], 3):
            fila['estado'] = 'Ausente'
    return filas


def escribir_hoja(ws, titulo, filas):
    ws.append([titulo])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append(HEADER)
    for celda in ws[2]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='1F4E78')
        celda.alignment = Alignment(horizontal='center')
    for f in filas:
        ws.append([f['n'], f['documento'], f['titular'], f['codigo'], f['sector'], f['zona'],
                   f['tipo'], f['rubro'], f['metraje'], f['estado'], f.get('fecha')])
    # El documento va como texto para conservar ceros a la izquierda.
    for fila in ws.iter_rows(min_row=3, min_col=2, max_col=2):
        fila[0].number_format = '@'
    for fila in ws.iter_rows(min_row=3, min_col=11, max_col=11):
        fila[0].number_format = 'yyyy-mm-dd'
    for col, ancho in zip('ABCDEFGHIJK', [5, 14, 34, 15, 22, 18, 17, 12, 11, 10, 12]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = 'A3'


def libro_valido(rng):
    wb = Workbook()
    wb.remove(wb.active)
    for hoja, filas in generar_filas(rng).items():
        escribir_hoja(wb.create_sheet(hoja),
                      f'PADRÓN DE COMERCIANTES - FERIA SALCEDO 2026 - {hoja} (DATOS DE PRUEBA)', filas)
    return wb


def libro_errores():
    wb = Workbook()
    ws = wb.active
    ws.title = 'BLOQUE C'
    base = dict(sector='Bloque C - Fila 1', zona='Zona Norte', tipo='Ferial', rubro='ropa', estado='Presente')
    casos = [
        # Fila válida de control.
        dict(documento='45871236', titular='CONTROL VALIDO QUISPE', codigo='ERR-C-001', metraje=3),
        dict(documento='', titular='SIN DOCUMENTO MAMANI', codigo='ERR-C-002', metraje=3),
        dict(documento='1234567', titular='DNI CORTO CONDORI', codigo='ERR-C-003', metraje=3),
        dict(documento='30123456789', titular='RUC INVALIDO TICONA', codigo='ERR-C-004', metraje=3),
        dict(documento='45871237', titular='', codigo='ERR-C-005', metraje=3),
        dict(documento='45871236', titular='DNI DUPLICADO CHOQUE', codigo='ERR-C-007', metraje=3),
        dict(documento='45871239', titular='CODIGO DUPLICADO HUANCA', codigo='ERR-C-001', metraje=3),
        dict(documento='45871240', titular='RUBRO INEXISTENTE PARI', codigo='ERR-C-009', metraje=3, rubro='ferreteria'),
        dict(documento='45871241', titular='AUSENTE SIN FECHA COAQUIRA', codigo='ERR-C-010', metraje=3, estado='Ausente'),
        dict(documento='45871242', titular='FECHA NO ES FERIA LAYME', codigo='ERR-C-011', metraje=3, fecha=date(2026, 10, 22)),
    ]
    filas = [{**base, **caso, 'n': i} for i, caso in enumerate(casos, start=1)]
    escribir_hoja(ws, 'PADRÓN DE PRUEBA CON ERRORES INTENCIONALES', filas)
    return wb


def main():
    OUT_DIR.mkdir(exist_ok=True)
    libro_valido(random.Random(2026)).save(OUT_DIR / 'padron-prueba-50.xlsx')
    libro_errores().save(OUT_DIR / 'padron-prueba-errores.xlsx')
    print(f'Libros generados en {OUT_DIR}')


if __name__ == '__main__':
    main()
